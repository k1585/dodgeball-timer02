# -*- coding: utf-8 -*-
"""만들어 둔 HTML 문서를 고칠 수 있는 문서 파일(.docx)로 바꾼다.

한글(한컴오피스)에서 바로 열리고 고칠 수 있다. 고친 뒤 '다른 이름으로 저장'에서
.hwp 를 고르면 한글 파일이 된다.

리브레오피스로 HTML을 바로 바꾸려 했으나 이 환경에는 HTML을 읽는 부분이 없어
실패했다. 그래서 HTML을 직접 읽어 문서를 짓는다. 이 HTML은 우리가 만든 것이라
쓰이는 표시가 몇 가지로 정해져 있어 그대로 옮길 수 있다.
"""
import base64, io, os, re, sys
from html.parser import HTMLParser
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, '한글문서')

NAVY = RGBColor(0x1E, 0x3A, 0x6B)
GREY = RGBColor(0x8A, 0x8F, 0xA3)
RED  = RGBColor(0xE4, 0x42, 0x3F)

# ---------- HTML 읽기 ----------

class Node:
    def __init__(self, tag, attrs=None):
        self.tag = tag
        self.attrs = dict(attrs or [])
        self.kids = []
        self.text = ''
    def cls(self):
        return self.attrs.get('class', '')

class Tree(HTMLParser):
    VOID = {'br', 'img', 'meta', 'link', 'hr'}
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.root = Node('root')
        self.stack = [self.root]
    def handle_starttag(self, tag, attrs):
        n = Node(tag, attrs)
        self.stack[-1].kids.append(n)
        if tag not in self.VOID:
            self.stack.append(n)
    def handle_endtag(self, tag):
        for i in range(len(self.stack) - 1, 0, -1):
            if self.stack[i].tag == tag:
                del self.stack[i:]
                return
    def handle_data(self, data):
        if data.strip():
            t = Node('#text')
            t.text = data
            self.stack[-1].kids.append(t)

def find(node, tag):
    if node.tag == tag:
        return node
    for k in node.kids:
        r = find(k, tag)
        if r:
            return r
    return None

# ---------- 글자 모양 ----------

def set_font(run, name='맑은 고딕', size=10, bold=False, color=None, mono=False):
    if mono:
        name = 'D2Coding'          # 없으면 한글이 알아서 비슷한 것으로 바꾼다
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts'); rpr.append(rf)
    for a in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'):
        rf.set(qn(a), name)

def shade(cell, hexcolor):
    el = OxmlElement('w:shd')
    el.set(qn('w:val'), 'clear')
    el.set(qn('w:color'), 'auto')
    el.set(qn('w:fill'), hexcolor)
    cell._tc.get_or_add_tcPr().append(el)

"""문단 속성(w:pPr)의 자식은 순서가 정해져 있다. 그냥 뒤에 붙이면 파일이
   규격에 어긋나 한글이나 워드가 열다가 고장 났다고 할 수 있다. 정해진 자리에 넣는다."""
PPR_ORDER = ['w:pStyle', 'w:keepNext', 'w:keepLines', 'w:pageBreakBefore',
  'w:framePr', 'w:widowControl', 'w:numPr', 'w:suppressLineNumbers', 'w:pBdr',
  'w:shd', 'w:tabs', 'w:suppressAutoHyphens', 'w:kinsoku', 'w:wordWrap',
  'w:overflowPunct', 'w:topLinePunct', 'w:autoSpaceDE', 'w:autoSpaceDN',
  'w:bidi', 'w:adjustRightInd', 'w:snapToGrid', 'w:spacing', 'w:ind',
  'w:contextualSpacing', 'w:mirrorIndents', 'w:suppressOverlap', 'w:jc',
  'w:textDirection', 'w:textAlignment', 'w:textboxTightWrap', 'w:outlineLvl',
  'w:divId', 'w:cnfStyle', 'w:rPr', 'w:sectPr', 'w:pPrChange']

def pPr_put(pPr, el, tag):
    after = PPR_ORDER[PPR_ORDER.index(tag) + 1:]
    pPr.insert_element_before(el, *after)

def para_shade(p, hexcolor):
    el = OxmlElement('w:shd')
    el.set(qn('w:val'), 'clear')
    el.set(qn('w:color'), 'auto')
    el.set(qn('w:fill'), hexcolor)
    pPr_put(p._p.get_or_add_pPr(), el, 'w:shd')

def para_border(p, hexcolor='1E3A6B', size=12, sides=('left',)):
    pPr = p._p.get_or_add_pPr()
    bd = OxmlElement('w:pBdr')
    for s in sides:
        e = OxmlElement('w:' + s)
        e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), str(size))
        e.set(qn('w:space'), '6'); e.set(qn('w:color'), hexcolor)
        bd.append(e)
    pPr_put(pPr, bd, 'w:pBdr')

# ---------- 줄 안의 글자들 ----------

def emit_inline(p, node, bold=False, mono=False, color=None, size=10):
    """<b>, <code>, <u>, <span class=sub/fill>, <br> 을 글자 모양으로 옮긴다"""
    for k in node.kids:
        if k.tag == '#text':
            t = re.sub(r'\s+', ' ', k.text)
            if not t.strip() and not p.runs:
                continue
            r = p.add_run(t)
            set_font(r, size=size, bold=bold, mono=mono, color=color)
        elif k.tag == 'br':
            p.add_run().add_break()
        elif k.tag in ('b', 'strong'):
            emit_inline(p, k, True, mono, color, size)
        elif k.tag == 'code':
            emit_inline(p, k, bold, True, color, size)
        elif k.tag == 'u':
            before = len(p.runs)
            emit_inline(p, k, bold, mono, color, size)
            for r in p.runs[before:]:
                r.font.underline = True
        elif k.tag == 'a':
            emit_inline(p, k, bold, mono, NAVY, size)
        elif k.tag == 'span':
            c = k.cls()
            if 'fill' in c:
                """채워 넣을 빈칸. 그냥 빈칸에 밑줄을 그으면, 그 빈칸이 줄 맨 끝에
                   올 때 밑줄이 그려지지 않는다(줄 끝 공백은 지워서 그린다).
                   밑줄 문자를 직접 쓰면 어디서든 그대로 보인다."""
                r = p.add_run('_' * 14)
                set_font(r, size=size, color=RGBColor(0x66, 0x6C, 0x7E))
            elif 'sub' in c:
                emit_inline(p, k, bold, mono, GREY, size - 0.5)
            elif 'no' in c:
                emit_inline(p, k, True, mono, RED, size)
            elif 'ok' in c:
                emit_inline(p, k, True, mono, RGBColor(0x1F, 0xA4, 0x55), size)
            else:
                emit_inline(p, k, bold, mono, color, size)
        else:
            emit_inline(p, k, bold, mono, color, size)

def plain(node):
    out = []
    def walk(n):
        if n.tag == '#text':
            out.append(n.text)
        for k in n.kids:
            walk(k)
    walk(node)
    return re.sub(r'\s+', ' ', ''.join(out)).strip()

# ---------- 덩어리별로 옮기기 ----------

def add_callout(doc, node, fill, bar):
    """box / key / warn — 배경색 있는 안내 상자"""
    first = True
    for part in split_block(node):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8 if first else 0)
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.left_indent = Cm(0.3)
        para_shade(p, fill)
        para_border(p, bar, 18, ('left',))
        if part.tag == 'li':
            r = p.add_run('· '); set_font(r, size=10)
            p.paragraph_format.left_indent = Cm(0.7)
        emit_inline(p, part)
        first = False

def split_block(node):
    """상자 안을 줄 단위로 쪼갠다. ul 이 있으면 li 하나가 한 줄."""
    out, buf = [], Node('span')
    for k in node.kids:
        if k.tag in ('ul', 'ol'):
            if buf.kids:
                out.append(buf); buf = Node('span')
            for li in k.kids:
                if li.tag == 'li':
                    out.append(li)
        else:
            buf.kids.append(k)
    if buf.kids:
        out.append(buf)
    return out or [node]

def add_table(doc, node):
    rows = [r for r in node.kids if r.tag == 'tr']
    if not rows:
        return
    ncol = 0
    for r in rows:
        n = 0
        for c in r.kids:
            if c.tag in ('th', 'td'):
                n += int(c.attrs.get('colspan', 1))
        ncol = max(ncol, n)
    t = doc.add_table(rows=len(rows), cols=ncol)
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = True

    """HTML 의 rowspan 은 '아래 칸까지 이어진다'는 뜻이다. 내용만 넣고 끝내면
       아래 칸이 빈 채로 남아 표가 이상해진다. 어디를 이어야 하는지 적어 두었다가
       내용을 다 넣은 뒤에 합친다. 합치는 도중에 칸 번호가 흔들리기 때문이다."""
    taken = set()          # (행, 열) — 위에서 이어져 내려온 자리
    spans = []             # 나중에 합칠 곳
    for ri, r in enumerate(rows):
        ci = 0
        for c in r.kids:
            if c.tag not in ('th', 'td'):
                continue
            while (ri, ci) in taken:
                ci += 1
            if ci >= ncol:
                break
            cell = t.cell(ri, ci)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            emit_inline(p, c, bold=(c.tag == 'th'), size=9.5)
            if c.tag == 'th':
                shade(cell, 'F2F5FB')
            rs = int(c.attrs.get('rowspan', 1))
            cs = int(c.attrs.get('colspan', 1))
            for rr in range(ri, min(ri + rs, len(rows))):
                for cc in range(ci, min(ci + cs, ncol)):
                    if (rr, cc) != (ri, ci):
                        taken.add((rr, cc))
            if rs > 1 or cs > 1:
                spans.append((ri, ci, min(ri + rs - 1, len(rows) - 1),
                                      min(ci + cs - 1, ncol - 1)))
            ci += cs
    for r0, c0, r1, c1 in spans:
        if (r0, c0) != (r1, c1):
            try:
                t.cell(r0, c0).merge(t.cell(r1, c1))
            except Exception:
                pass

    """칸 너비 — HTML 에 적어 둔 폭이 있으면 그 비율대로 나눈다.
       없으면 고르게 나눈다. 표 폭과 칸 폭을 둘 다 적어야 한글에서도 그대로 나온다."""
    total = Cm(16.6)
    hints = []
    if rows:
        ci = 0
        for c in rows[0].kids:
            if c.tag not in ('th', 'td'):
                continue
            m = re.search(r'width:\s*(\d+)px', c.attrs.get('style', ''))
            hints.append(int(m.group(1)) if m else None)
    if len(hints) == ncol and any(h for h in hints):
        known = sum(h for h in hints if h)
        blanks = [i for i, h in enumerate(hints) if not h]
        if blanks:
            rest = max(total.cm - known * total.cm / max(known + 260, 1), total.cm * 0.3)
            each = rest / len(blanks)
            widths = [(h * total.cm / (known + 260)) if h else each for h in hints]
        else:
            widths = [h * total.cm / known for h in hints]
    else:
        widths = [total.cm / ncol] * ncol
    scale = total.cm / sum(widths)
    widths = [w * scale for w in widths]
    t.columns  # noqa — 폭은 칸마다 적어야 워드·한글이 따른다
    for ri in range(len(rows)):
        for ci in range(ncol):
            try:
                t.cell(ri, ci).width = Cm(widths[ci])
            except Exception:
                pass
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def add_shots(doc, node, tmpdir, counter):
    """화면 사진 여러 장을 한 줄에 놓는다 (표를 써서 나란히)"""
    shots = []
    def walk(n):
        if n.tag == 'div' and 'shot' in n.cls():
            img = find(n, 'img')
            cap = None
            for k in n.kids:
                if k.tag == 'div' and 'cap' in k.cls():
                    cap = plain(k)
            if img is not None:
                shots.append((img.attrs.get('src', ''), cap, n))
            return
        for k in n.kids:
            walk(k)
    walk(node)
    if not shots:
        return
    per = min(len(shots), 3)
    t = doc.add_table(rows=0, cols=per)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i in range(0, len(shots), per):
        chunk = shots[i:i + per]
        cells = t.add_row().cells
        for j, (src, cap, holder) in enumerate(chunk):
            m = re.match(r'data:image/(\w+);base64,(.*)', src, re.S)
            if not m:
                continue
            counter[0] += 1
            ext = 'jpg' if m.group(1) in ('jpeg', 'jpg') else m.group(1)
            path = os.path.join(tmpdir, 'shot%03d.%s' % (counter[0], ext))
            with open(path, 'wb') as f:
                f.write(base64.b64decode(m.group(2)))
            cell = cells[j]
            cell.text = ''
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(path, width=Cm(4.4))
            if cap:
                cp = cell.add_paragraph()
                cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = cp.add_run(cap)
                set_font(r, size=8.5, color=GREY)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

# ---------- 문서 하나 만들기 ----------

def convert(src, dst, tmpdir):
    with open(src, encoding='utf-8') as f:
        html = f.read()
    # <style> 안쪽은 글이 아니다
    html = re.sub(r'<style.*?</style>', '', html, flags=re.S)
    tr = Tree(); tr.feed(html)
    body = find(tr.root, 'body') or tr.root

    doc = Document()
    st = doc.styles['Normal']
    st.font.name = '맑은 고딕'
    st.font.size = Pt(10)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), '맑은 고딕')
    for s in doc.sections:
        s.top_margin = Cm(2.0); s.bottom_margin = Cm(2.0)
        s.left_margin = Cm(2.2); s.right_margin = Cm(2.2)

    counter = [0]
    for node in body.kids:
        tag, c = node.tag, node.cls()
        if tag == '#text':
            continue
        if tag == 'div' and 'doc-kind' in c:
            p = doc.add_paragraph()
            r = p.add_run('[ %s ]' % plain(node))
            set_font(r, size=9, bold=True, color=NAVY)
            p.paragraph_format.space_after = Pt(2)
        elif tag == 'h1':
            p = doc.add_paragraph()
            r = p.add_run(plain(node))
            set_font(r, size=19, bold=True, color=NAVY)
            p.paragraph_format.space_after = Pt(2)
        elif tag == 'p' and 'updated' in c:
            p = doc.add_paragraph()
            r = p.add_run(plain(node))
            set_font(r, size=9, color=GREY)
            para_border(p, 'DCE3EE', 6, ('bottom',))
            p.paragraph_format.space_after = Pt(14)
        elif tag == 'h2':
            p = doc.add_paragraph()
            r = p.add_run(plain(node))
            set_font(r, size=13, bold=True, color=NAVY)
            para_border(p, 'DCE3EE', 8, ('bottom',))
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(6)
        elif tag == 'h3':
            p = doc.add_paragraph()
            r = p.add_run(plain(node))
            set_font(r, size=11, bold=True, color=RGBColor(0x33, 0x43, 0x6B))
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(3)
        elif tag == 'div' and 'key' in c:
            add_callout(doc, node, 'FFF8E8', 'E0A800')
        elif tag == 'div' and 'warn' in c:
            add_callout(doc, node, 'FDF0EF', 'E4423F')
        elif tag == 'div' and 'box' in c:
            add_callout(doc, node, 'F7FAFF', '9FB5DA')
        elif tag == 'div' and 'step' in c:
            p = doc.add_paragraph()
            para_border(p, 'C9D6EE', 18, ('left',))
            p.paragraph_format.left_indent = Cm(0.4)
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            emit_inline(p, node)
        elif tag in ('p', 'div') and 'cmd' in c:
            p = doc.add_paragraph()
            para_shade(p, 'F1F4FA')
            p.paragraph_format.left_indent = Cm(0.4)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(6)
            emit_inline(p, node, mono=True, size=9)
        elif tag == 'div' and ('row' in c or 'shot' in c):
            add_shots(doc, node, tmpdir, counter)
        elif tag == 'table':
            add_table(doc, node)
        elif tag in ('ul', 'ol'):
            for i, li in enumerate([k for k in node.kids if k.tag == 'li'], 1):
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.7)
                p.paragraph_format.space_after = Pt(2)
                r = p.add_run(('%d. ' % i) if tag == 'ol' else '· ')
                set_font(r, size=10)
                emit_inline(p, li)
        elif tag == 'p':
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            if 'sub' in c:
                emit_inline(p, node, color=GREY, size=9.5)
            else:
                emit_inline(p, node)
        elif tag == 'footer':
            p = doc.add_paragraph()
            para_border(p, 'DCE3EE', 6, ('top',))
            p.paragraph_format.space_before = Pt(18)
            r = p.add_run(plain(node))
            set_font(r, size=8.5, color=GREY)
    # 기본 서식에 들어 있는 zoom 에 값이 빠져 있어 규격에 어긋난다. 채워 준다.
    z = doc.settings.element.find(qn('w:zoom'))
    if z is not None and z.get(qn('w:percent')) is None:
        z.set(qn('w:percent'), '100')
    doc.save(dst)
    return os.path.getsize(dst)


def main():
    import tempfile, glob
    os.makedirs(OUT, exist_ok=True)
    srcs = sorted(glob.glob(os.path.join(HERE, '[0-9][0-9]_*.html')))
    if not srcs:
        raise SystemExit('바꿀 HTML이 없습니다')
    print('고칠 수 있는 문서로 바꿉니다\n')
    with tempfile.TemporaryDirectory() as tmp:
        for s in srcs:
            name = os.path.splitext(os.path.basename(s))[0]
            dst = os.path.join(OUT, name + '.docx')
            size = convert(s, dst, tmp)
            print('  ✓ %-40s %6.0f KB' % (name + '.docx', size / 1024))
    # 글로만 된 것도 같이 옮긴다
    import shutil
    for t in glob.glob(os.path.join(HERE, '[0-9][0-9]_*.txt')):
        shutil.copy(t, OUT)
        print('  ✓ %-40s (그대로)' % os.path.basename(t))
    print('\n끝났습니다. 한글에서 열어 고치시면 됩니다.')

if __name__ == '__main__':
    main()
